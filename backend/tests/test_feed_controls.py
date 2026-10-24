from io import BytesIO, StringIO
import csv
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from openpyxl import load_workbook
from pypdf import PdfReader

from app.db import Base, get_db
from app.entitlements import ENTITLEMENTS
from app.routers import events
from app.services import feed_controls as controls
from test_events_filters import _event, _stub_enrichment


@pytest.fixture
def client(monkeypatch):
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    db = Session(engine)
    db.add_all([_event(i, "congress_trade", symbol="AAPL", member_name="Test member", trade_type="purchase",
                      payload={"report_date": "2026-10-01"}) for i in range(1, 1002)])
    db.commit()
    _stub_enrichment(monkeypatch)
    monkeypatch.setattr(events, "_enqueue_missing_trade_outcomes", lambda *a, **k: None)
    monkeypatch.setattr(controls, "current_user", lambda db, request, **k: request.cookies.get("test_tier"))
    monkeypatch.setattr(controls, "entitlements_for_user", lambda db, tier: ENTITLEMENTS[tier])
    events._EVENTS_RESPONSE_CACHE.clear()
    app = FastAPI()
    app.include_router(events.router, prefix="/api")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[events.rate_limit_provider_backed] = lambda: None
    with TestClient(app) as test_client:
        yield test_client
    db.close()
    engine.dispose()


@pytest.mark.parametrize("tier,cap", [(None, 25), ("free", 25), ("premium", 1000), ("pro", 1000), ("admin", 1000)])
@pytest.mark.parametrize("alias", ["limit", "page_size"])
def test_feed_size_cap_is_enforced_before_query_and_cache(client, tier, cap, alias):
    if tier:
        client.cookies.set("test_tier", tier)
    response = client.get(f"/api/events?{alias}=1000&enrich_prices=false", headers={"X-CT-Entitlement-Tier": "pro"})
    assert response.status_code == 200, response.text
    assert len(response.json()["items"]) == cap
    assert response.json()["limit"] == cap
    assert response.json()["has_more"] is True


@pytest.mark.parametrize("tier,allowed", [(None, False), ("free", False), ("premium", False), ("pro", True), ("admin", True)])
def test_sort_gate_ignores_spoofed_tier_header(client, tier, allowed):
    if tier:
        client.cookies.set("test_tier", tier)
    response = client.get("/api/events?sort_by=amount&sort_dir=asc&enrich_prices=false",
                          headers={"X-CT-Entitlement-Tier": "pro"})
    assert response.status_code == (200 if allowed else 403), response.text


ROW = ["17", "Oct 1, 2026", "2026-09-29", "AAPL", "A & B Inc.", "René Test", "Congress",
       "Purchase", "$1K - $15K", "Pending", "-", "https://example.com/disclosure"]


@pytest.mark.parametrize("tier,allowed", [(None, False), ("free", False), ("premium", False), ("pro", True), ("admin", True)])
@pytest.mark.parametrize("extension", ["csv", "xlsx", "pdf"])
def test_export_gate_all_formats(client, tier, allowed, extension):
    if tier:
        client.cookies.set("test_tier", tier)
    response = client.post(f"/api/events/export.{extension}", json={"rows": [ROW]},
                           headers={"X-CT-Entitlement-Tier": "pro"})
    assert response.status_code == (200 if allowed else 403), response.text if not allowed else ""
    if allowed:
        assert response.headers["cache-control"] == "private, no-store"
        assert "attachment;" in response.headers["content-disposition"]


def test_export_limits_and_shape(client):
    client.cookies.set("test_tier", "pro")
    for rows in ([], [ROW] * 1001, [["bad"]], [[*ROW[:-1], "x" * 2049]]):
        assert client.post("/api/events/export.csv", json={"rows": rows}).status_code == 422
    assert client.post("/api/events/export.exe", json={"rows": [ROW]}).status_code == 422


def test_paid_page_cannot_be_reused_after_downgrade(client):
    client.cookies.set("test_tier", "premium")
    query = "/api/events?limit=1000&page_size=1000&enrich_prices=false"
    assert len(client.get(query).json()["items"]) == 1000
    client.cookies.set("test_tier", "free")
    response = client.get(query)
    assert response.status_code == 200
    assert len(response.json()["items"]) == 25
    assert response.json()["limit"] == 25


def test_csv_excel_round_trip_and_formula_safety():
    rows = [ROW, [*ROW[:4], '=HYPERLINK("https://evil.test")', *ROW[5:]]]
    content, _ = controls.render_feed_export(rows, "csv")
    parsed = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
    assert parsed[1][4:6] == ROW[4:6]
    assert parsed[2][4].startswith("'=")
    content, _ = controls.render_feed_export(rows, "xlsx")
    book = load_workbook(BytesIO(content))
    sheet = book.active
    assert sheet["E3"].data_type == "s"
    assert sheet["E3"].value == rows[1][4]
    assert sheet["F2"].value == "René Test"
    assert sheet["J2"].value == "Pending"
    assert sheet["L2"].value == ROW[11]
    assert sheet.freeze_panes == "A2"


def test_pdf_and_excel_retain_all_1000_rows():
    rows = [[str(i), *ROW[1:]] for i in range(1, 1001)]
    xlsx, _ = controls.render_feed_export(rows, "xlsx")
    assert load_workbook(BytesIO(xlsx)).active.max_row == 1001
    pdf, _ = controls.render_feed_export(rows, "pdf")
    document = PdfReader(BytesIO(pdf))
    assert len(document.pages) > 1
    text = "\n".join(page.extract_text() for page in document.pages)
    assert text.count("René Test") == 1000
    assert "1000" in text
    assert "Disclosure URL" in text
    assert "Pending" in text
