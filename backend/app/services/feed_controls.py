"""Feed access policy and bounded exports of the user's displayed page."""
from __future__ import annotations

import csv
import re
from io import BytesIO, StringIO
from typing import Annotated
from xml.sax.saxutils import escape

from fastapi import HTTPException, Request
from pydantic import BaseModel, Field, StringConstraints
from sqlalchemy.orm import Session
from app.auth import current_user
from app.entitlements import entitlements_for_user

FEED_MAX_ROWS = 1000
EXPORT_COLUMNS = ("Event ID", "Reported / filing date", "Transaction date", "Ticker",
                  "Company", "Person / Entity", "Source", "Action", "Amount",
                  "% Delta", "Signal", "Disclosure URL")
Cell = Annotated[str, StringConstraints(max_length=2048)]
ExportRow = Annotated[list[Cell], Field(min_length=12, max_length=12)]


class FeedExportRequest(BaseModel):
    rows: list[ExportRow] = Field(min_length=1, max_length=FEED_MAX_ROWS)


def feed_rank(db: Session, request: Request | None) -> int:
    user = current_user(db, request, required=False) if request is not None and hasattr(request, "cookies") else None
    return entitlements_for_user(db, user).rank if user is not None else 0


def require_feed_pro(request: Request, db: Session) -> None:
    if feed_rank(db, request) < 20:
        raise HTTPException(403, "Pro is required for feed sorting and exports.")


def feed_limit(db: Session, request: Request | None, limit: int, page_size: int | None,
               sort_by: str | None, sort_dir: str | None) -> int:
    rank = feed_rank(db, request)
    if (sort_by or sort_dir) and rank < 20:
        raise HTTPException(403, "Pro is required for feed sorting and exports.")
    return min(max(page_size if page_size is not None else limit, 1), FEED_MAX_ROWS if rank >= 10 else 25)


def render_feed_export(rows: list[list[str]], export_format: str) -> tuple[bytes, str]:
    rows = [[re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value) for value in row] for row in rows]
    if export_format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(EXPORT_COLUMNS)
        for row in rows:
            # CSV quoting alone does not prevent spreadsheet formula execution.
            writer.writerow(["'" + v if v.lstrip().startswith(("=", "+", "-", "@")) or v.startswith(("\t", "\r", "\n")) else v for v in row])
        return output.getvalue().encode("utf-8-sig"), "text/csv"
    if export_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = "Feed"
        sheet.append(EXPORT_COLUMNS)
        for row in rows:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="164E3F")
        for index, width in enumerate((12, 24, 20, 12, 40, 36, 20, 24, 22, 20, 12, 65), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output = BytesIO()
        book.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if export_format != "pdf":
        raise ValueError("Unsupported feed export format")

    from pathlib import Path
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

    if "FeedVera" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("FeedVera", str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")))
    style = ParagraphStyle("Feed", fontName="FeedVera", fontSize=7, leading=10, wordWrap="CJK")
    heading = ParagraphStyle("FeedHeading", parent=style, fontSize=16, leading=22)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=28, bottomMargin=28,
                            title="Walnut Markets feed", author="Walnut Markets")
    widths = [30, 66, 60, 44, 85, 85, 56, 60, 65, 47, 38, 118]
    data = [[Paragraph(escape(value), style) for value in EXPORT_COLUMNS]]
    data.extend([[Paragraph(escape(value).replace("\n", "<br/>"), style) for value in row] for row in rows])
    class FeedTable(LongTable):
        def split(self, available_width, available_height):
            # Keep ordinary rows together. Only split an oversized row when it
            # cannot fit on an otherwise empty page.
            self.splitInRow = 0
            parts = super().split(available_width, available_height)
            if not parts and available_height >= 480:
                self.splitInRow = 1
                parts = super().split(available_width, available_height)
            return parts

    table = FeedTable(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D6EDE5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F5")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    def footer(canvas, document):
        canvas.setFont("FeedVera", 8)
        canvas.drawRightString(818, 14, f"Page {document.page}")

    doc.build([Paragraph("Walnut Markets | Feed", heading),
               Paragraph(f"Current displayed page: {len(rows)} rows, in the selected order. Dates and missing values are retained as displayed.", style),
               Spacer(1, 12), table], onFirstPage=footer, onLaterPages=footer)
    return output.getvalue(), "application/pdf"
